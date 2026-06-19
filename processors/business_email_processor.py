from services.report_email_service import ReportEmailService
from openpyxl import load_workbook
from datetime import datetime
import os
import json
import logging

class BusinessEmailProcess:

    def run(self):

        excel_template = "shared/download/daily_business_report.xlsx"
        
        download_folder = (
            f"shared/download/Daily Business Report "
            f"{datetime.now().strftime('%m-%d-%Y')}"
        )

        with open("shared/Json/file.json", "r") as file:
            data = json.load(file)

        wb = load_workbook(excel_template)
        ws = wb.active

        for item in data["File"]:

            target_keyword = item["file_name"]
            target_cell = item["cell"]

            target_file = None

            for file in os.listdir(download_folder):

                if target_keyword in file:
                    target_file = file
                    break

            if target_file:

                logging.info(f"Detected: {target_file}")

                ws[target_cell] = target_file

            else:

                logging.error(f"File not found: {target_keyword}")

                ws[target_cell] = "No Report Generated"


        save_folder = (
            f"shared/download/Daily Business Report "
            f"{datetime.now().strftime('%m-%d-%Y')}"
        )

        os.makedirs(save_folder, exist_ok=True)


        file_path = os.path.join(
            save_folder,
            f"Daily_Business_Report-{datetime.now().strftime('%m-%d-%Y')}.xlsx"
        )

        wb.save(file_path)

    
        logging.info("Excel updated successfully!")

        # SEND EMAIL
        report = ReportEmailService()
        report.send_email()