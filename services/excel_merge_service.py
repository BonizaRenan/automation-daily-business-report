from openpyxl import load_workbook, Workbook
import logging
import os
import csv


class ExcelMergeService:

    # =========================
    # MERGE FILES
    # =========================
    def merge_files(self, file_list, output_file, config_item=None):

        priority_rules = (config_item or {}).get("priority_rules", {})

        def get_priority(file_path):

            name = os.path.basename(file_path).lower()

            for key, priority in priority_rules.items():

                if key.lower() in name:
                    return priority

            return 99

        # SORT FILES
        file_list = sorted(file_list, key=get_priority)

        print("\nFILES TO MERGE:")
        for f in file_list:
            print(f)

        new_wb = Workbook()
        new_wb.remove(new_wb.active)

        for file in file_list:

            try:

                extension = os.path.splitext(file)[1].lower()

                sheet_name = os.path.splitext(
                    os.path.basename(file)
                )[0][:31]

                new_ws = new_wb.create_sheet(title=sheet_name)

                # =========================
                # CSV FILE
                # =========================
                if extension == ".csv":

                    print(f"READING CSV: {file}")

                    with open(
                        file,
                        mode="r",
                        encoding="utf-8-sig",
                        newline=""
                    ) as csv_file:

                        reader = csv.reader(csv_file)

                        for row in reader:
                            new_ws.append(row)

                # =========================
                # EXCEL FILE
                # =========================
                else:

                    print(f"READING EXCEL: {file}")

                    wb = load_workbook(file)
                    ws = wb.active

                    for row in ws.iter_rows(values_only=True):
                        new_ws.append(row)

                logging.info(f"Merged: {file}")

            except Exception as e:

                print(f"ERROR MERGING {file}: {e}")

                logging.error(
                    f"Failed merging file {file}: {e}"
                )

        # =========================
        # SAVE FILE
        # =========================
        new_wb.save(output_file)

        logging.info(f"Merged file created: {output_file}")