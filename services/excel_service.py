from openpyxl import load_workbook, Workbook
from datetime import datetime
import logging
import pandas as pd
import os


class ExcelService:

    # =========================
    # CLEAN VALUES
    # =========================
    def clean_value(self, value):
        if value in ([], None):
            return ""

        if isinstance(value, list):
            return ", ".join(map(str, value))

        if isinstance(value, dict):
            return str(value)

        return value

    # =========================
    # EXPORT
    # =========================
    def export(self, data, output_file=None):

        logging.info("Exporting results to Excel")

        df = pd.DataFrame(data)

        if output_file is None:
            today = datetime.now().strftime("%m-%d-%Y")
            folder_name = f"Daily Business Report {today}"
            output_folder = os.path.join("shared", "download", folder_name)
            os.makedirs(output_folder, exist_ok=True)

            output_file = os.path.join(output_folder, "Export_Report.xlsm")

        template_path = "shared/download/template.xlsm"

        wb = load_workbook(template_path, keep_vba=True)
        ws = wb.active

        ws.delete_rows(1, ws.max_row)

        # headers
        for col_index, column in enumerate(df.columns, start=1):
            ws.cell(row=1, column=col_index, value=column)

        # rows
        for row_index, row in enumerate(df.itertuples(index=False), start=2):
            for col_index, value in enumerate(row, start=1):
                ws.cell(
                    row=row_index,
                    column=col_index,
                    value=self.clean_value(value)
                )

        wb.save(output_file)

        logging.info(f"Export completed: {output_file}")

    # =========================
    # GENERIC MERGE ENGINE
    # =========================
    def merge_files(self, file_list, output_file, config_item=None):

        priority_rules = (config_item or {}).get("priority_rules", {})

        def get_priority(file_path):
            name = os.path.basename(file_path).lower()

            for key, priority in priority_rules.items():
                if key.lower() in name:
                    return priority

            return 99

        file_list = sorted(file_list, key=get_priority)

        new_wb = Workbook()
        new_wb.remove(new_wb.active)

        for file in file_list:
            try:
                wb = load_workbook(file)
                ws = wb.active

                sheet_name = os.path.splitext(os.path.basename(file))[0][:31]
                new_ws = new_wb.create_sheet(title=sheet_name)

                for row in ws.iter_rows(values_only=True):
                    new_ws.append(row)

            except Exception as e:
                logging.error(f"Failed merging file {file}: {e}")

        new_wb.save(output_file)

        logging.info(f"Merged file created: {output_file}")